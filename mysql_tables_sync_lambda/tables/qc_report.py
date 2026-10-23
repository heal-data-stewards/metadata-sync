"""
Builds the QC Report -- a styled, multi-sheet .xlsx summarizing data
quality across progress_tracker, nihtables, mysql_merged, and
study_lookup_table, written to S3.

Mirrors mysql_code/python_scripts/HEAL_07_QC.py exactly (including its own
numbering -- section "2. awards table" is in the original script's header
comment but was never actually implemented there either; this port
doesn't invent it). The original's near-identical Count/Percent/Total
tab-building code (repeated for 1A/1B/1C/1D/3A/3C/4B/5B) is consolidated
into one _status_table() helper -- same computed values, no repeated
boilerplate.

All four inputs (progress_tracker, nihtables, mysql_merged,
study_lookup_table) are either pre-existing real MySQL tables or already
built by this lambda -- no manual input, no new MySQL table (this is a
report, not a table).
"""

import io
import logging
import os
import re
from datetime import datetime, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from db import clean_mysql_value

logger = logging.getLogger(__name__)

QC_REPORT_FILE = "QCReport.xlsx"

_FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="1B365D")
_FONT_SUBTITLE = Font(name="Calibri", size=12, bold=True)
_FONT_BOLD = Font(name="Calibri", size=11, bold=True)
_FONT_REGULAR = Font(name="Calibri", size=11)
_FILL_HEADER = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")


def _read_table(conn, table: str) -> pd.DataFrame:
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM `{table}`")
    cols = [d[0] for d in cur.description]
    df = pd.DataFrame(cur.fetchall(), columns=cols)
    cur.close()
    for col in df.columns:
        df[col] = df[col].map(clean_mysql_value)
    if "appl_id" in df.columns:
        df["appl_id"] = df["appl_id"].astype(str)
    return df


# ---------------------------------------------------------------------------
# Sheet-building helpers
# ---------------------------------------------------------------------------

def _add_header(ws, text, size=14):
    ws.append([])
    row = ws.max_row + 1
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _FONT_TITLE if size == 14 else _FONT_SUBTITLE


def _add_text(ws, text):
    row = ws.max_row + 1
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _FONT_REGULAR


def _clean_row(values):
    return [None if pd.isna(v) else v for v in values]


def _add_table(ws, df: pd.DataFrame, title: str | None = None):
    if title:
        row = ws.max_row + 1
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = _FONT_BOLD

    if df.empty:
        ws.append(["No records found matching criteria."])
        ws.cell(row=ws.max_row, column=1).font = _FONT_REGULAR
        ws.append([])
        return

    ws.append(list(df.columns))
    header_row = ws.max_row
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = _FONT_BOLD
        cell.fill = _FILL_HEADER

    for _, row_vals in df.iterrows():
        ws.append(_clean_row(row_vals.values))
        data_row = ws.max_row
        for c in range(1, len(df.columns) + 1):
            ws.cell(row=data_row, column=c).font = _FONT_REGULAR

    ws.append([])


def _dump_df_to_new_sheet(wb: Workbook, df: pd.DataFrame, sheet_name: str):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(list(df.columns))
    for _, row_vals in df.iterrows():
        ws.append(_clean_row(row_vals.values))


def _status_table(series: pd.Series, status_map: dict | None = None) -> pd.DataFrame:
    """Replicates the Count/Percent/Total tab repeated throughout the
    original script for a single flag/status column."""
    counts = series.value_counts(dropna=False)
    percentages = (series.value_counts(dropna=False, normalize=True) * 100).round(1)
    tab = pd.DataFrame({"Status": counts.index, "Count": counts.values, "Percent": percentages.values})
    if status_map:
        tab["Status"] = tab["Status"].map(status_map).fillna(tab["Status"])
    total_row = pd.DataFrame([{"Status": "Total", "Count": tab["Count"].sum(), "Percent": tab["Percent"].sum().round(1)}])
    return pd.concat([tab, total_row], ignore_index=True)


# ---------------------------------------------------------------------------
# Build logic
# ---------------------------------------------------------------------------

def build_qc_report(
    progress_tracker_df: pd.DataFrame,
    nihtables_df: pd.DataFrame,
    mysql_merged_df: pd.DataFrame,
    study_lookup_df: pd.DataFrame,
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "QC Report"

    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    _add_header(ws, "--------------QC Report--------------", size=20)
    _add_header(ws, f"--------------Run Date: {today} --------------")

    # --- 1. progress_tracker table ---
    _add_header(ws, "")
    _add_header(ws, "--------------1. progress_tracker table--------------")
    _add_text(ws, "The MYSQL progress_tracker table is automatically updated daily with a pull of fresh data from the Platform MDS.")

    pt_df = progress_tracker_df.copy()

    pt_df["1a_miss_appl_id"] = (pt_df["appl_id"].isna() | (pt_df["appl_id"].astype(str) == "0")).astype(int)
    tab_1a = _status_table(pt_df["1a_miss_appl_id"], {0: "0. Not Missing", 1: "1. Missing"})
    _add_table(ws, tab_1a, "1A. Number of records missing appl_id")
    pt_df.insert(0, "1a_miss_appl_id", pt_df.pop("1a_miss_appl_id"))

    pt_df["project_num_str"] = pt_df["project_num"].fillna("").astype(str).str.strip()
    pt_df["1b_miss_proj_num"] = (pt_df["project_num_str"] == "").astype(int)
    pt_df["num_dashes"] = pt_df["project_num_str"].apply(lambda x: x.count("-"))
    pt_df["ctn_flag"] = pt_df["project_num_str"].str.contains(r"^CTN", flags=re.IGNORECASE).astype(int)
    pt_df["1c_fmt_proj_num"] = ((pt_df["num_dashes"] > 1) | (pt_df["ctn_flag"] == 1)).astype(int)

    tab_1b = _status_table(pt_df["1b_miss_proj_num"], {0: "0. Not Missing", 1: "1. Missing"})
    _add_header(ws, "")
    _add_table(ws, tab_1b, "1B. Number of records missing project_num")
    pt_df.insert(1, "1b_miss_proj_num", pt_df.pop("1b_miss_proj_num"))

    tab_1c = _status_table(pt_df["1c_fmt_proj_num"], {0: "0. Good Format", 1: "1. Bad Format"})
    _add_header(ws, "")
    _add_table(ws, tab_1c, "1C. Number of records with bad format for project_num")
    pt_df.insert(2, "1c_fmt_proj_num", pt_df.pop("1c_fmt_proj_num"))

    pt_df["1d_fmt_ctn_flag"] = ((pt_df["1c_fmt_proj_num"] == 1) & (pt_df["ctn_flag"] == 1)).astype(int)
    bad_proj_num_df = pt_df[pt_df["1c_fmt_proj_num"] == 1]
    tab_1d = _status_table(bad_proj_num_df["1d_fmt_ctn_flag"], {0: "0. Bad Format not due to CTN Flag", 1: "1. Bad Format Due to CTN Flag"})
    _add_header(ws, "")
    _add_table(ws, tab_1d, "1D. Number of bad project_number formats due to CTN protocol in project_num field")
    pt_df.insert(3, "1d_fmt_ctn_flag", pt_df.pop("1d_fmt_ctn_flag"))

    # --- 3. Compare appl_ids in MySQL tables ---
    _add_header(ws, "")
    _add_header(ws, "")
    _add_header(ws, "--------------3. Compare appl_ids in MySQL tables--------------")

    nih_tables = nihtables_df.copy()
    nih_tables["3a_merge_reporter_awards"] = nih_tables["merge_reporter_awards"]
    tab_3a = _status_table(nih_tables["3a_merge_reporter_awards"])
    _add_header(ws, "")
    _add_table(ws, tab_3a, "3A. Compare: reporter and awards (using merge_reporter_awards from nih_tables)")
    nih_tables.insert(1, "3a_merge_reporter_awards", nih_tables.pop("3a_merge_reporter_awards"))

    mysql_df = mysql_merged_df.copy()
    mysql_df = mysql_df[mysql_df["entity_type"] != "CTN"]
    mysql_df["3c_merge_awards_mds"] = mysql_df["merge_awards_mds"]
    tab_3c = _status_table(mysql_df["3c_merge_awards_mds"])
    _add_header(ws, "")
    _add_table(ws, tab_3c, "3C. Compare: MySQL [reporter & awards] and MDS (using merge_awards_mds from mysql)")
    mysql_df.insert(1, "3c_merge_awards_mds", mysql_df.pop("3c_merge_awards_mds"))

    # --- 4. Metrics by Study ---
    _add_header(ws, "")
    _add_header(ws, "")
    _add_header(ws, "--------------4. Metrics by Study--------------")

    study_lookup = study_lookup_df.copy()
    xstudies = study_lookup["xstudy_id"].dropna().drop_duplicates().astype(int)
    summary_stats = pd.DataFrame([{"Statistic": "xstudy_id", "Distinct Count": len(xstudies)}])
    _add_header(ws, "")
    _add_table(ws, summary_stats, "4A. Number of studies, by Stewards xstudy_id")

    hdp_metrics = study_lookup[["xstudy_id", "study_hdp_id"]].drop_duplicates().copy()
    hdp_metrics["4b_study_has_hdp"] = (hdp_metrics["study_hdp_id"].notna() & (hdp_metrics["study_hdp_id"] != "")).astype(int)
    tab_4b = _status_table(hdp_metrics["4b_study_has_hdp"], {0: "0. No HDP_ID", 1: "1. HDP_ID"})
    _add_header(ws, "")
    _add_table(ws, tab_4b, "4B. Number of studies with/out a HDP ID")
    hdp_metrics.insert(0, "4b_study_has_hdp", hdp_metrics.pop("4b_study_has_hdp"))

    # --- 5. Entity Types ---
    _add_header(ws, "")
    _add_header(ws, "")
    _add_header(ws, "--------------5. Entity Types--------------")
    _add_header(ws, "--------------5A. Entity Type Rules--------------", size=12)
    _add_text(ws,
        "There are 3 entity types in our data: Study, CTN Protocol & Other. Every record is assigned Study for "
        "entity_type by default. The entity type is changed to CTN or Other if it meets one of the following conditions:")
    _add_text(ws, "- replace entity_type=CTN if the project_num field in Platform/progress_tracker table data is a CTN Protocol number")
    _add_text(ws, "- replace entity_type=CTN if res_net==CTN")
    _add_text(ws,
        "- replace entity_type=Other if the project_num field in Platform/progress_tracker table data does not match "
        "the default project number format for grants. For example an OT contract like the Stewards would be replaced entity_type='Other'")
    _add_text(ws, "- replace entity_type=Other if appl_id==0")

    _add_header(ws, "--------------5B. Entity Type: Other--------------", size=12)
    mysql_reports = mysql_merged_df.copy()
    tab_5b = _status_table(mysql_reports["entity_type"])
    _add_header(ws, "")
    _add_table(ws, tab_5b, "5B. Entity Type")

    other_entities = mysql_reports[mysql_reports["entity_type"] == "Other"][
        ["entity_type", "appl_id", "proj_num", "hdp_id", "archived"]
    ].copy()
    other_entities["5b_entity_type"] = other_entities["entity_type"]
    other_entities.insert(0, "5b_entity_type", other_entities.pop("5b_entity_type"))

    # --- Raw-data dump sheets, same as the original's appended ExcelWriter sheets ---
    _dump_df_to_new_sheet(wb, pt_df, "1_progress_tracker")
    _dump_df_to_new_sheet(wb, nih_tables.drop(columns=["proj_abs"], errors="ignore"), "3a_nih_tables")
    _dump_df_to_new_sheet(wb, mysql_df.drop(columns=["proj_abs"], errors="ignore"), "3c_mysql")
    _dump_df_to_new_sheet(wb, study_lookup, "4a_study_lookup")
    _dump_df_to_new_sheet(wb, hdp_metrics, "4b_hdp_metrics")
    _dump_df_to_new_sheet(wb, other_entities, "5b_other_entities")

    logger.info("QC report built: %d distinct studies, %d 'Other' entities", len(xstudies), len(other_entities))
    return wb


def _to_xlsx_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def update_qc_report(
    conn,
    ref_source,
    nihtables_conn=None,
    mysql_merged_conn=None,
    study_lookup_table_conn=None,
    nihtables_table: str | None = None,
    mysql_merged_table: str | None = None,
    study_lookup_table_name: str | None = None,
    report_file: str | None = None,
) -> dict:
    """
    conn: connection for reading progress_tracker -- a genuinely external
          input this lambda doesn't build, always on DB_NAME.
    nihtables_conn / mysql_merged_conn / study_lookup_table_conn:
          connections for reading those three tables, which this lambda
          builds. nihtables/mysql_merged ALWAYS resolve to tempspace;
          study_lookup_table follows DB_ENVIRONMENT. Default to `conn` if
          not given.
    ref_source: a sources.FileSource pointed at the directory/prefix to
          write the report into.
    """
    nihtables_conn = nihtables_conn or conn
    mysql_merged_conn = mysql_merged_conn or conn
    study_lookup_table_conn = study_lookup_table_conn or conn

    if nihtables_table is None:
        nihtables_table = os.getenv("NIHTABLES_TABLE_NAME", "nihtables_test")
    if mysql_merged_table is None:
        mysql_merged_table = os.getenv("MYSQL_MERGED_TABLE_NAME", "mysql_merged_test")
    if study_lookup_table_name is None:
        study_lookup_table_name = os.getenv("STUDY_LOOKUP_TABLE_NAME", "study_lookup_table")
    if report_file is None:
        report_file = os.getenv("QC_REPORT_FILE", QC_REPORT_FILE)

    progress_tracker_df = _read_table(conn, "progress_tracker")
    nihtables_df = _read_table(nihtables_conn, nihtables_table)
    mysql_merged_df = _read_table(mysql_merged_conn, mysql_merged_table)
    study_lookup_df = _read_table(study_lookup_table_conn, study_lookup_table_name)

    wb = build_qc_report(progress_tracker_df, nihtables_df, mysql_merged_df, study_lookup_df)
    ref_source.write_bytes(report_file, _to_xlsx_bytes(wb))

    return {"report": report_file}
