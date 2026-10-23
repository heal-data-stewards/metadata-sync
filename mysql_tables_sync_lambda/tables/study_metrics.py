"""
Builds the Study Metrics report -- fills a hand-maintained Excel template
(uploaded to S3, like research_networks' lookup workbook) with HDE study
metrics computed from mysql_merged and study_lookup_table, plus two
raw-data drill-down sheets, and writes the filled workbook back to S3 as
a new report.

Mirrors mysql_code/python_scripts/HEAL_09_StudyMetrics.py, with two fixes:

1. The original hardcodes wb["8. Repository Names"], but the template's
   actual sheet name is "7. Repository Names" -- a KeyError that silently
   truncates every real run partway through step 7. Confirmed against the
   real Sept 2026 production output: everything from the repo-breakdown
   sheet onward (VLMD-on-Platform count, data-linked count, the CTN/MySQL
   study/MySQL award counts, and both drill-down sheets) is blank there.
   This port uses the template's real sheet name and runs to completion.

2. The main "metrics" base dataset used merge_awards_mds != "In MDS only",
   which is backwards. Cross-checked against mysql_code/stata_scripts/
   HEAL_09_StudyMetrics.do (the Stata original this script was itself
   ported from): "drop if merge_awards_mds==1 /* keep only records that
   appear in Platform */". Stata's merge code 1 is the master-only side;
   tables/mysql_merge.py's own indicator=... map confirms the master
   (left) side is "In MySQL only" (merge_awards_mds's left_only ->
   "In MySQL only", right_only -> "In MDS only"). So the original script
   drops "In MySQL only" to keep Platform-registered records -- the
   Python port (and this port, until this fix) instead dropped
   "In MDS only", which is the inverse. CTN protocols, ICPSR, Intramural,
   and most "Other" entities are registered on the Platform but have no
   matching NIH award in MySQL's awards table, so they're almost entirely
   "In MDS only" -- the inverted filter silently excluded nearly all of
   them, undercounting total HEAL studies and zeroing out the CTN/Other
   breakdown. The CTN-protocol-count step (C76) and the MySQL-awards
   steps (C80/C83-85) are unaffected -- see build_study_metrics for which
   base each one actually uses, matching the Stata original's own three
   different bases (metrics: drop "In MySQL only"; CTN count: no filter
   at all; awards: drop "In MDS only", Stata's merge code 2).

Metric #9 ("studies with VLMD available in HSS") is not computed here,
same as the original -- its own comment says to see a separate report.
"""

import io
import logging
import os

import pandas as pd
from openpyxl import load_workbook

from db import clean_mysql_value

logger = logging.getLogger(__name__)

STUDY_METRICS_TEMPLATE_NAME = "StudyMetrics_template_v2.xlsx"
STUDY_METRICS_REPORT_FILE = "StudyMetrics.xlsx"

_REPO_SHEET = "7. Repository Names"


def _read_table(conn, table: str) -> pd.DataFrame:
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM `{table}`")
    cols = [d[0] for d in cur.description]
    df = pd.DataFrame(cur.fetchall(), columns=cols)
    cur.close()
    for col in df.columns:
        df[col] = df[col].map(clean_mysql_value)
    if "appl_id" in df.columns:
        df["appl_id"] = df["appl_id"].astype(str)
    return df


def _write_drill_sheet(wb, sheet_name: str, df: pd.DataFrame, autofilter: bool = False, freeze: str | None = None) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    if autofilter:
        ws.auto_filter.ref = ws.dimensions
    if freeze:
        ws.freeze_panes = freeze


# ---------------------------------------------------------------------------
# Build logic
# ---------------------------------------------------------------------------

def build_study_metrics(
    mysql_merged_df: pd.DataFrame,
    study_lookup_df: pd.DataFrame,
    template_bytes: bytes,
    run_date: str,
) -> bytes:
    # Platform-registered records (drop "In MySQL only" -- see module
    # docstring fix #2): the base for every "studies" metric below.
    metrics_base = mysql_merged_df[mysql_merged_df["merge_awards_mds"] != "In MySQL only"].copy()
    metrics = metrics_base[metrics_base["guid_type"].isin(["discovery_metadata", "unregistered_discovery_metadata"])].copy()
    logger.info("study_metrics: metrics_base=%d rows, metrics=%d rows", len(metrics_base), len(metrics))

    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb["Metrics"]
    ws["C2"] = run_date

    # --- 1. Number of HEAL Studies ---
    ws["C5"] = len(metrics)
    ws["C9"] = int((metrics["entity_type"] == "Study").sum())
    ws["C10"] = int((metrics["entity_type"] == "CTN").sum())
    ws["C11"] = int((metrics["entity_type"] == "Other").sum())

    # --- 2. Number of studies registered ---
    ws["C18"] = int((metrics["is_registered"] == "registered").sum())
    ws["C19"] = int((metrics["is_registered"] == "not registered").sum())

    # --- 3. Studies who've submitted CDE usage ---
    num_cde = pd.to_numeric(metrics["num_common_data_elements"], errors="coerce")
    ws["C25"] = int((num_cde > 0).sum())

    # --- 4. Studies submitted SLMD ---
    pct_complete = pd.to_numeric(metrics["overall_percent_complete"], errors="coerce")
    slmd = pct_complete >= 50
    ws["C31"] = int(slmd.sum())
    ws["C32"] = int((~slmd).sum())

    # --- 5. HEAL studies by data sharing intention ---
    ws["C39"] = int(metrics["gen3_data_availability"].isna().sum())
    ws["C40"] = int((metrics["gen3_data_availability"] == "unaccessible").sum())
    ws["C41"] = int((metrics["gen3_data_availability"] == "not_available").sum())
    ws["C42"] = int((metrics["gen3_data_availability"] == "mixed_availability").sum())

    # --- 6. Number of studies selected repo ---
    producing_sharing = metrics[metrics["gen3_data_availability"] != "not_available"]
    ws["C50"] = int((producing_sharing["repository_selected"] == "Yes").sum())
    ws["C51"] = int((producing_sharing["repository_selected"] == "No").sum())

    # --- 7. Number of studies selecting each repo (own worksheet) ---
    # Matches the original: a blank/NaN repository_name passes this filter
    # (NaN != "" is True) and only becomes "Missing" via the fillna below.
    with_repo = producing_sharing[producing_sharing["repository_name"].str.strip() != ""].copy()
    with_repo["repository_name"] = with_repo["repository_name"].fillna("Missing")
    repo_counts = with_repo["repository_name"].value_counts(dropna=False).reset_index()
    repo_counts.columns = ["repository_name", "count"]
    repo_counts["percent"] = (repo_counts["count"] / repo_counts["count"].sum() * 100).round(2)
    repo_counts["cum_percent"] = repo_counts["percent"].cumsum().round(2)
    total_row = pd.DataFrame([{
        "repository_name": "Total",
        "count": repo_counts["count"].sum(),
        "percent": round(repo_counts["percent"].sum(), 2),
        "cum_percent": None,
    }])
    repo_counts = pd.concat([repo_counts, total_row], ignore_index=True)

    repo_ws = wb[_REPO_SHEET]
    for row_idx, row_data in enumerate(repo_counts.itertuples(index=False)):
        for col_idx, value in enumerate(row_data):
            repo_ws.cell(row=3 + row_idx, column=1 + col_idx).value = value

    # --- 8. Studies with VLMD on Platform ---
    num_dd = pd.to_numeric(metrics["num_data_dictionaries"], errors="coerce")
    ws["C59"] = int((num_dd > 0).sum())

    # --- 9. Studies with VLMD available in HSS: not computed here, same as original ---

    # --- 10. Studies with data linked on Platform ---
    ws["C69"] = int((producing_sharing["data_linked_on_platform"] == "Yes").sum())
    ws["C70"] = int((producing_sharing["data_linked_on_platform"] == "No").sum())

    # --- 11. Other: MySQL entities and awards ---
    # CTN protocol count: no merge_awards_mds filter in the Stata original
    # at all -- just mds_ctn_flag, against the full unfiltered table.
    ctn = mysql_merged_df[mysql_merged_df["mds_ctn_flag"] == "1"]
    ws["C76"] = len(ctn)

    xstudy_id = pd.to_numeric(study_lookup_df["xstudy_id"], errors="coerce")
    ws["C78"] = int(xstudy_id.nunique())

    # Number of MySQL awards / awards by entity type: a DIFFERENT base
    # than `metrics` above -- Stata drops "In MDS only" here (merge code
    # 2), the opposite of the "In MySQL only" drop metrics_base uses,
    # since an "award" by definition comes from the MySQL/awards side.
    awards_base = mysql_merged_df[mysql_merged_df["merge_awards_mds"] != "In MDS only"]
    awards = awards_base[awards_base["appl_id"] != ""][["appl_id", "entity_type"]].drop_duplicates(subset="appl_id")
    ws["C80"] = len(awards)
    ws["C83"] = int((awards["entity_type"] == "Study").sum())
    ws["C84"] = int((awards["entity_type"] == "CTN").sum())
    ws["C85"] = int((awards["entity_type"] == "Other").sum())

    # --- Drill-down sheets, for manual QC of the numbers above ---
    mysql_drill = mysql_merged_df.drop(columns=["proj_abs"], errors="ignore")
    target_cols = ["compound_key", "appl_id", "hdp_id", "merge_awards_mds", "merge_reporter_awards"]
    cols = [c for c in target_cols if c in mysql_drill.columns] + [c for c in mysql_drill.columns if c not in target_cols]
    mysql_drill = mysql_drill[cols].sort_values(by="compound_key")

    _write_drill_sheet(wb, "mysql_today", mysql_drill, autofilter=True, freeze="F2")
    _write_drill_sheet(wb, "study_lookup_table", study_lookup_df)

    logger.info("study_metrics: workbook built, sheets=%s", wb.sheetnames)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def update_study_metrics(
    mysql_merged_conn,
    study_lookup_table_conn,
    manual_inputs_source,
    report_source,
    mysql_merged_table: str | None = None,
    study_lookup_table_name: str | None = None,
    template_name: str | None = None,
    report_file: str | None = None,
) -> dict:
    """
    mysql_merged_conn: connection for reading `mysql_merged` -- ALWAYS
          tempspace regardless of DB_ENVIRONMENT.
    study_lookup_table_conn: connection for reading `study_lookup_table`,
          which follows DB_ENVIRONMENT.
    manual_inputs_source: a sources.FileSource pointed at the manual-input
          prefix/directory the template workbook is uploaded to.
    report_source: a sources.FileSource pointed at the directory/prefix
          to write the filled report into.
    """
    if mysql_merged_table is None:
        mysql_merged_table = os.getenv("MYSQL_MERGED_TABLE_NAME", "mysql_merged_test")
    if study_lookup_table_name is None:
        study_lookup_table_name = os.getenv("STUDY_LOOKUP_TABLE_NAME", "study_lookup_table")
    if template_name is None:
        template_name = os.getenv("STUDY_METRICS_TEMPLATE_NAME", STUDY_METRICS_TEMPLATE_NAME)
    if report_file is None:
        report_file = os.getenv("STUDY_METRICS_REPORT_FILE", STUDY_METRICS_REPORT_FILE)

    mysql_merged_df = _read_table(mysql_merged_conn, mysql_merged_table)
    study_lookup_df = _read_table(study_lookup_table_conn, study_lookup_table_name)
    template_bytes = manual_inputs_source.read_bytes(template_name)

    from datetime import datetime, timezone
    run_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    data = build_study_metrics(mysql_merged_df, study_lookup_df, template_bytes, run_date)
    report_source.write_bytes(report_file, data)

    return {"report": report_file, "run_date": run_date}
